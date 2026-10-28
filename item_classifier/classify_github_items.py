#!/usr/bin/env python3
"""
classify_github_items.py

Read a tabular file (CSV, TSV, or XLSX) containing links to GitHub issues,
pull requests, or discussions; fetch each item's title and body from GitHub;
classify it by audience/task; and write a tabular file with:

    Title | Link | Classification

If an item's body states who the feature is intended for (for example, an
answer to "What kind of user is the feature intended for?" in an issue
template, or a line such as "Intended user: curators"), that statement is
extracted and given heavy weight in the classification.

The item's title and labels (e.g. "Feature: Account & User Info") are also
weighed: the title counts double in the keyword rules, and each label is
matched against the LABEL_RULES patterns below.

Classifications:
    - Data Depositors, Seekers, and Downloaders
    - Data Curators and Installation Administrators
    - Other

Usage:
    python classify_github_items.py input.csv -o output.csv
    python classify_github_items.py input.xlsx -o output.xlsx --link-column "URL"
    python classify_github_items.py input.csv --use-claude      # LLM classification

Environment variables:
    GITHUB_TOKEN        Recommended (higher rate limits; required for discussions)
    ANTHROPIC_API_KEY   Only needed with --use-claude

Dependencies:
    requests            (pip install requests)
    openpyxl            (only for .xlsx input/output)
"""

import argparse
import csv
import json
import os
import re
import sys
import time

import requests

# --------------------------------------------------------------------------- #
# Categories
# --------------------------------------------------------------------------- #
DEPOSITORS = "Data Depositors, Seekers, and Downloaders"
CURATORS = "Data Curators and Installation Administrators"
OTHER = "Other"
FETCH_ERROR = "Unclassified (could not fetch item)"

CATEGORIES = [DEPOSITORS, CURATORS, OTHER]

# --------------------------------------------------------------------------- #
# Rubric -> keyword rules (edit these to tune the classifier)
# Each entry is (regex, weight). Matching is case-insensitive.
# --------------------------------------------------------------------------- #
RULES = {
    # Common end-user tasks
    DEPOSITORS: [
        (r"\b(create|creating|register|registering|sign[- ]?up|sign[- ]?in|log[- ]?in)\b.{0,30}\baccount", 3),
        (r"\b(search|searching|browse|browsing|find(ing)?)\b.{0,30}\b(data|datasets?|files?|collections?)\b", 2),
        (r"\brequest(ing)?\b.{0,20}\b(access|permission)\b", 3),
        (r"\brestricted (files?|data)\b", 2),
        (r"\bdownload(s|ing|ed)?\b", 2),
        (r"\b(deposit(ing|or|ors)?|upload(ing)?|submit(ting)?)\b.{0,30}\b(data|datasets?|files?)\b", 3),
        (r"\b(create|creating|add|adding)\b.{0,20}\bdatasets?\b", 2),
        (r"\bedit(ing)?\b.{0,20}\b(datasets?|metadata|files?)\b", 2),
        (r"\bterms of (use|access)\b|\blicen[sc]e\b|\bterms\b.{0,10}\bdatasets?\b", 2),
        (r"\bguest ?books?\b", 3),
        (r"\b(end[- ]?users?|researchers?|depositors?)\b", 1),
        (r"\bdataset (page|landing page)\b", 1),
    ],
    # Common data curation and installation administration tasks
    CURATORS: [
        (r"\breview(ing)?\b.{0,20}\bdatasets?\b|\bdatasets?\b.{0,20}\breview\b", 3),
        (r"\bpublish(ing)?\b.{0,20}\b(datasets?|versions?)\b", 3),
        (r"\b(submit(ted)? for review|return(ed)? to (author|contributor)|in review)\b", 3),
        (r"\b(add|adding|invite|inviting|remove|removing)\b.{0,30}\b(members?|users?|groups?)\b.{0,30}\b(collections?|dataverses?)\b", 3),
        (r"\b(collection|dataverse)\b.{0,20}\b(members?|roles?|permissions?)\b", 3),
        (r"\b(change|changing|assign|assigning|grant|granting|manage|managing)\b.{0,20}\b(permissions?|roles?)\b", 2),
        (r"\bcuration (status|labels?|workflow|tasks?)\b|\bcurat(or|ors|ion|ing)\b", 3),
        (r"\b(contributor|curator|admin) role\b", 2),
        (r"\bpublication workflow\b|\bpublishing workflow\b", 2),
        # Installation-level administration of content and users
        (r"\bcustom metadata blocks?\b", 4),
        (r"\bmetadata blocks?\b.{0,30}\b(tsv|load(ing)?|install(ing)?|add(ing)?)\b", 3),
        (r"\b(superuser|super[- ]user)\b", 3),
        (r"\b(admin(istrator)?|superuser) dashboard\b", 3),
        (r"\b(deactivat(e|ing)|merg(e|ing)|delet(e|ing)|manag(e|ing))\b.{0,20}\b(user|account)s?\b", 2),
        (r"\binstallation[- ](wide|level)\b|\bnetwork[- ]wide\b", 2),
        (r"\binstallation admin(istrators?)?\b|\bsysadmins?\b", 3),
    ],
    # System configuration, deployment, development, etc.
    OTHER: [
        (r"\b(configur(e|ation|ing)|jvm[- ]?options?|database settings?|:[A-Z][A-Za-z]+ setting)\b", 2),
        (r"\b(install(ation|ing)?|deploy(ment|ing)?|upgrad(e|ing)|docker|kubernetes|payara|solr|postgres(ql)?)\b", 2),
        (r"\b(build|maven|ci|pipeline|unit tests?|integration tests?|refactor(ing)?|dependency|dependencies)\b", 2),
        (r"\b(api (endpoint|token)|rest api|native api|sword|oai-pmh|harvest(ing)?)\b", 1),
        (r"\b(documentation|docs|release notes?)\b", 1),
        (r"\b(s3|storage driver|dns|doi provider|datacite|mail server|smtp)\b", 2),
    ],
}

COMPILED = {
    cat: [(re.compile(pattern, re.IGNORECASE | re.DOTALL), w) for pattern, w in rules]
    for cat, rules in RULES.items()
}

MIN_SCORE = 3          # below this, the item is classified as "Other"
TITLE_MULTIPLIER = 2   # title matches count extra


def strip_noise(text: str) -> str:
    """Remove fenced code blocks, inline code, URLs and HTML comments."""
    text = re.sub(r"```.*?```", " ", text, flags=re.DOTALL)
    text = re.sub(r"<!--.*?-->", " ", text, flags=re.DOTALL)
    text = re.sub(r"`[^`]*`", " ", text)
    text = re.sub(r"https?://\S+", " ", text)
    return text


# --------------------------------------------------------------------------- #
# Explicit "intended user" detection
# --------------------------------------------------------------------------- #
# Matches questions/labels such as:
#   "What kind of user is the feature intended for?"   "Who is this for?"
#   "Intended user:"   "Target audience"   "User type / persona"
QUESTION_RE = re.compile(
    r"(what|which)\s+(kind|type|sort)s?\s+of\s+(users?|people|audience)"
    r"|(what|which)\s+users?\b"
    r"|who\s+(is|are)\s+(this|the\s+feature|the\s+change|it|these)\b.{0,40}\b(for|intended)"
    r"|who\s+(will|would|should)\s+(use|benefit)"
    r"|(intended|target)\s+(users?|audience|persona|personas|user\s+types?)"
    r"|\buser\s+(type|role|persona)s?\b|\bpersonas?\b",
    re.IGNORECASE,
)
BOLD_ONLY_RE = re.compile(r"^(\*\*|__).+(\*\*|__)\s*:?\s*$")
PLACEHOLDER_RE = re.compile(
    r"^(no response|n/?a|none|tbd|tba|unknown|not sure|\?+|-+|\.+)$", re.IGNORECASE
)

# How the author's stated user maps onto the three categories.
USER_RULES = {
    DEPOSITORS: re.compile(
        r"\b(depositors?|data (seekers?|downloaders?|consumers?|users?|reusers?)|seekers?|"
        r"downloaders?|end[- ]?users?|regular users?|general users?|researchers?|scientists?|"
        r"students?|authors?|public|anonymous|guests?|authenticated users?|"
        r"logged[- ]in users?|dataset (owners?|creators?))\b",
        re.IGNORECASE,
    ),
    CURATORS: re.compile(
        r"\b(curators?|installation admin(istrator)?s?|(collection|dataverse|repository|"
        r"content) (admin(istrator)?s?|managers?|owners?)|admins?|administrators?|"
        r"super[- ]?users?|sysadmins?|system admin(istrator)?s?|librarians?|data stewards?|"
        r"reviewers?|moderators?)\b",
        re.IGNORECASE,
    ),
    OTHER: re.compile(
        r"\b(developers?|devs?|integrators?|system integrators?|devops|operators?|"
        r"maintainers?|software engineers?|programmers?|api (users?|clients?)|"
        r"tool builders?|third[- ]party (developers?|apps?|tools?)|deployers?)\b",
        re.IGNORECASE,
    ),
}
EXPLICIT_USER_BONUS = 10  # added to every category the author's answer matches


def _clean_md(line: str) -> str:
    line = re.sub(r"^\s{0,3}#{1,6}\s*", "", line)
    line = re.sub(r"[*_`>]+", "", line)
    return line.strip()


def extract_intended_user(body: str):
    """Return the author's stated intended user(s) as text, or None.

    Handles issue-form style headings ("### What kind of user ...?" followed by
    the answer), bold labels, and inline "Label: answer" lines. Ignores empty
    answers and unchecked checkboxes.
    """
    body = re.sub(r"<!--.*?-->", " ", body or "", flags=re.DOTALL)
    lines = body.splitlines()

    for i, raw in enumerate(lines):
        clean = _clean_md(raw)
        m = QUESTION_RE.search(clean)
        if not m or len(clean) > 300:
            continue

        stripped = raw.strip()
        is_heading = stripped.startswith("#") or bool(BOLD_ONLY_RE.match(stripped))
        inline = ""
        sep = re.search(r"[?:]\s*(\S.*)$", clean[m.end():])
        if sep:
            inline = sep.group(1).strip()
        if not (is_heading or clean.endswith(("?", ":")) or inline):
            continue  # an ordinary sentence that happens to mention users

        parts = [inline] if inline else []
        for nxt in ([] if inline else lines[i + 1:i + 16]):  # inline answer is complete
            s_nxt = nxt.strip()
            if s_nxt.startswith("#") or BOLD_ONLY_RE.match(s_nxt):
                break  # next section / next question
            if not s_nxt:
                if parts:
                    break  # end of the answer
                continue
            if s_nxt.endswith("?") and len(s_nxt) < 200:
                break  # next question
            parts.append(s_nxt)

        # If the answer is a checkbox list, keep only the checked options.
        if any(re.match(r"^\s*[-*]\s*\[[ xX]\]", p) for p in parts):
            parts = [re.sub(r"^\s*[-*]\s*\[[xX]\]\s*", "", p)
                     for p in parts if re.match(r"^\s*[-*]\s*\[[xX]\]", p)]

        answer = " ".join(_clean_md(p) for p in parts).strip()
        if answer and not PLACEHOLDER_RE.match(answer):
            return answer[:500]
    return None


# --------------------------------------------------------------------------- #
# Label handling
# --------------------------------------------------------------------------- #
LABEL_WEIGHT = 4        # added once per label per category it matches
LABEL_USER_BONUS = 6    # for labels that name a user type, e.g. "User: Curator"

# Matched (case-insensitively) against each label name. Edit to fit your repo's
# label taxonomy, e.g. "Feature: Account & User Info", "Feature: Search/Browse".
LABEL_RULES = {
    DEPOSITORS: re.compile(
        r"\b(account|user info|log[- ]?in|sign[- ]?(up|in)|search|browse|download\w*|"
        r"upload\w*|deposit\w*|guest ?books?|terms of (use|access)|licen[sc]es?|"
        r"request(ing)? access|restricted|dataset page|landing page|"
        r"edit(ing)? (datasets?|metadata)|end[- ]?users?)\b",
        re.IGNORECASE,
    ),
    CURATORS: re.compile(
        r"\b(curat\w*|review\w*|publish\w*|permissions?|roles?|metadata blocks?|"
        r"custom metadata|collection (members?|permissions?|admin\w*)|"
        r"(installation|dataverse|collection) admin\w*|superusers?|moderat\w*)\b",
        re.IGNORECASE,
    ),
    OTHER: re.compile(
        r"\b(config\w*|install\w*|deploy\w*|docker|kubernetes|solr|postgres\w*|payara|"
        r"api|sword|oai|harvest\w*|build|ci|release|tests?|testing|infrastructure|"
        r"performance|security|integrations?|documentation|docs|dependenc\w+|upgrade|"
        r"storage|s3|doi|dev(elopment)?|tech(nical)? debt|refactor\w*|scalab\w+|"
        r"monitoring|logging)\b",
        re.IGNORECASE,
    ),
}

# Labels that explicitly name a user type, e.g. "User: Curator", "Persona: Developer"
USER_LABEL_RE = re.compile(
    r"^\s*(user(\s*(type|role|persona))?|role|persona|audience|intended\s+user|target\s+user)"
    r"\s*[:/\-]\s*(?P<value>.+)$",
    re.IGNORECASE,
)


def score_labels(labels, scores):
    """Add label-based evidence to `scores` in place."""
    for label in labels or []:
        m = USER_LABEL_RE.match(label)
        if m:  # explicit user-type label
            for cat, rx in USER_RULES.items():
                if rx.search(m.group("value")):
                    scores[cat] += LABEL_USER_BONUS
            continue
        for cat, rx in LABEL_RULES.items():
            if rx.search(label):
                scores[cat] += LABEL_WEIGHT


def classify_by_keywords(title: str, body: str, labels=None):
    """Return (classification, intended_user_text_or_None)."""
    intended = extract_intended_user(body or "")
    title, body = strip_noise(title or ""), strip_noise(body or "")

    scores = {cat: 0 for cat in CATEGORIES}
    for cat, rules in COMPILED.items():
        for rx, weight in rules:
            scores[cat] += weight * len(rx.findall(body))
            scores[cat] += weight * TITLE_MULTIPLIER * len(rx.findall(title))

    # Labels (e.g. "Feature: Account & User Info") are a strong, curated signal.
    score_labels(labels, scores)

    # The author's explicit statement of the intended user is strong evidence.
    if intended:
        for cat, rx in USER_RULES.items():
            if rx.search(intended):
                scores[cat] += EXPLICIT_USER_BONUS

    best_cat = max(CATEGORIES, key=lambda c: (scores[c], -CATEGORIES.index(c)))
    if scores[best_cat] < MIN_SCORE:
        return OTHER, intended
    return best_cat, intended


# --------------------------------------------------------------------------- #
# Optional LLM classification (uses the rubric directly)
# --------------------------------------------------------------------------- #
RUBRIC_PROMPT = f"""You classify GitHub issues/pull requests/discussions for a data repository platform.
Choose exactly one of three categories based on the primary task or audience the item concerns.

1. "{DEPOSITORS}": common end-user tasks: creating an account, searching for data,
   requesting access to datasets, downloading data, depositing datasets, editing
   datasets, adding datasets' terms of use, creating guestbooks.
2. "{CURATORS}": common data curation and installation administration tasks:
   reviewing datasets, publishing datasets, adding new collection members, changing
   collection user permissions, changing curation status, adding custom metadata blocks.
3. "{OTHER}": any other tasks, including system configuration and administration tasks
   (deployment, server settings, development work, etc.).

If the item states who the feature is intended for (for example, an answer to
"What kind of user is the feature intended for?"), that statement is given to you
as "Author's stated intended user". Treat it as strong evidence and let it guide
your choice unless the rest of the item clearly contradicts it.

The item's labels are also provided. Labels such as "Feature: Account & User Info",
"Feature: Search/Browse" or "Feature: File Upload & Handling" point to end-user tasks;
labels about curation, review, publishing, permissions or metadata blocks point to
curators/installation administrators; labels about configuration, installation,
deployment, APIs or development point to "Other". Use the title and labels together
with the body.

Respond with ONLY the category name, exactly as written above."""


def classify_with_claude(title: str, body: str, model: str, intended=None, labels=None) -> str:
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        sys.exit("ANTHROPIC_API_KEY is not set (required for --use-claude).")
    resp = requests.post(
        "https://api.anthropic.com/v1/messages",
        headers={
            "x-api-key": api_key,
            "anthropic-version": "2023-06-01",
            "content-type": "application/json",
        },
        json={
            "model": model,
            "max_tokens": 50,
            "system": RUBRIC_PROMPT,
            "messages": [{
                "role": "user",
                "content": (
                    f"Title: {title}\n\n"
                    f"Labels: {', '.join(labels) if labels else '(none)'}\n\n"
                    f"Author's stated intended user: {intended or '(not stated)'}\n\n"
                    f"Body:\n{(body or '')[:8000]}"
                ),
            }],
        },
        timeout=60,
    )
    resp.raise_for_status()
    answer = "".join(
        b.get("text", "") for b in resp.json().get("content", []) if b.get("type") == "text"
    ).strip().strip('"')
    for cat in CATEGORIES:
        if answer.lower() == cat.lower():
            return cat
    for cat in CATEGORIES:  # tolerate extra words around the category name
        if cat.lower() in answer.lower():
            return cat
    return OTHER


# --------------------------------------------------------------------------- #
# GitHub access
# --------------------------------------------------------------------------- #
GITHUB_URL_RE = re.compile(
    r"https?://(?:www\.)?github\.com/([^/\s]+)/([^/\s]+)/(issues|pull|discussions)/(\d+)",
    re.IGNORECASE,
)


def parse_github_url(url: str):
    m = GITHUB_URL_RE.search(url or "")
    if not m:
        return None
    owner, repo, kind, number = m.groups()
    return owner, repo, kind.lower(), int(number)


def github_headers():
    headers = {"Accept": "application/vnd.github+json", "User-Agent": "classify-github-items"}
    token = os.environ.get("GITHUB_TOKEN")
    if token:
        headers["Authorization"] = f"Bearer {token}"
    return headers


def github_request(method, url, **kwargs):
    """Request with simple rate-limit handling and retries."""
    for attempt in range(4):
        resp = requests.request(method, url, headers=github_headers(), timeout=30, **kwargs)
        if resp.status_code in (403, 429) and resp.headers.get("X-RateLimit-Remaining") == "0":
            reset = int(resp.headers.get("X-RateLimit-Reset", time.time() + 60))
            wait = max(reset - int(time.time()), 1) + 1
            print(f"  Rate limited; sleeping {wait}s...", file=sys.stderr)
            time.sleep(min(wait, 3700))
            continue
        if resp.status_code >= 500:
            time.sleep(2 ** attempt)
            continue
        return resp
    return resp


def fetch_item(owner, repo, kind, number):
    """Return (title, body, labels) for an issue, PR, or discussion."""
    if kind in ("issues", "pull"):
        # The issues endpoint also serves pull requests.
        resp = github_request("GET", f"https://api.github.com/repos/{owner}/{repo}/issues/{number}")
        resp.raise_for_status()
        data = resp.json()
        labels = [l["name"] if isinstance(l, dict) else str(l) for l in data.get("labels", [])]
        return data.get("title", ""), data.get("body") or "", labels

    # Discussions are only available through GraphQL, which requires a token.
    if not os.environ.get("GITHUB_TOKEN"):
        raise RuntimeError("GITHUB_TOKEN is required to read discussions")
    query = """
    query($owner:String!, $repo:String!, $number:Int!) {
      repository(owner:$owner, name:$repo) {
        discussion(number:$number) { title body labels(first: 50) { nodes { name } } }
      }
    }"""
    resp = github_request(
        "POST", "https://api.github.com/graphql",
        json={"query": query, "variables": {"owner": owner, "repo": repo, "number": number}},
    )
    resp.raise_for_status()
    payload = resp.json()
    disc = ((payload.get("data") or {}).get("repository") or {}).get("discussion")
    if not disc:
        raise RuntimeError(f"Discussion not found: {payload.get('errors')}")
    labels = [n["name"] for n in ((disc.get("labels") or {}).get("nodes") or [])]
    return disc.get("title", ""), disc.get("body") or "", labels


# --------------------------------------------------------------------------- #
# Tabular I/O
# --------------------------------------------------------------------------- #
def read_table(path):
    """Return (header, rows) where rows is a list of lists of strings."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError:
            sys.exit("Reading .xlsx requires openpyxl: pip install openpyxl")
        ws = openpyxl.load_workbook(path, read_only=True, data_only=True).active
        data = [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
    else:
        delimiter = "\t" if ext in (".tsv", ".tab") else ","
        with open(path, newline="", encoding="utf-8-sig") as f:
            data = list(csv.reader(f, delimiter=delimiter))
    if not data:
        sys.exit("Input file is empty.")
    return data[0], data[1:]


def write_table(path, header, rows):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError:
            sys.exit("Writing .xlsx requires openpyxl: pip install openpyxl")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(header)
        for row in rows:
            ws.append(row)
        wb.save(path)
    else:
        delimiter = "\t" if ext in (".tsv", ".tab") else ","
        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f, delimiter=delimiter)
            writer.writerow(header)
            writer.writerows(rows)


def find_link_columns(header, rows, requested):
    """Return indices of columns that hold GitHub links."""
    if requested:
        lowered = [h.strip().lower() for h in header]
        if requested.strip().lower() not in lowered:
            sys.exit(f"Column '{requested}' not found. Columns: {header}")
        return [lowered.index(requested.strip().lower())]
    hits = [
        i for i in range(len(header))
        if any(i < len(r) and GITHUB_URL_RE.search(r[i] or "") for r in rows)
    ]
    if not hits:
        sys.exit("No column containing GitHub issue/PR/discussion links was found. "
                 "Use --link-column to name it.")
    return hits[:1]  # use the first matching column


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input", help="Input CSV/TSV/XLSX with GitHub links")
    ap.add_argument("-o", "--output", help="Output file (default: <input>_classified.<ext>)")
    ap.add_argument("--link-column", help="Name of the column containing GitHub links (auto-detected if omitted)")
    ap.add_argument("--use-claude", action="store_true",
                    help="Classify with the Anthropic API instead of keyword rules")
    ap.add_argument("--model", default="claude-sonnet-5-5", help="Model for --use-claude")
    ap.add_argument("--delay", type=float, default=0.2, help="Seconds to pause between GitHub requests")
    args = ap.parse_args()

    header, rows = read_table(args.input)
    col = find_link_columns(header, rows, args.link_column)[0]

    base, ext = os.path.splitext(args.input)
    output = args.output or f"{base}_classified{ext or '.csv'}"

    results, cache = [], {}
    links = [(r[col].strip() if col < len(r) else "") for r in rows]
    links = [l for l in links if l]

    for n, link in enumerate(links, 1):
        if link in cache:
            results.append(cache[link])
            continue
        parsed = parse_github_url(link)
        if not parsed:
            print(f"[{n}/{len(links)}] Skipping non-GitHub-item link: {link}", file=sys.stderr)
            row = ["", link, FETCH_ERROR]
        else:
            try:
                title, body, labels = fetch_item(*parsed)
                intended = extract_intended_user(body)
                if args.use_claude:
                    label = classify_with_claude(title, body, args.model, intended, labels)
                else:
                    label, _ = classify_by_keywords(title, body, labels)
                row = [title, link, label]
                note = ""
                if labels:
                    note += f" [labels: {', '.join(labels)[:80]}]"
                if intended:
                    note += f" [intended user: {intended[:60]}]"
                print(f"[{n}/{len(links)}] {label}: {title[:70]}{note}", file=sys.stderr)
            except Exception as exc:  # keep going; record the failure
                print(f"[{n}/{len(links)}] ERROR {link}: {exc}", file=sys.stderr)
                row = ["", link, FETCH_ERROR]
            time.sleep(args.delay)
        cache[link] = row
        results.append(row)

    write_table(output, ["Title", "Link", "Classification"], results)
    print(f"\nWrote {len(results)} rows to {output}", file=sys.stderr)


if __name__ == "__main__":
    main()
